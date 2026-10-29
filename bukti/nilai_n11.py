#!/usr/bin/env python3
"""
nilai_n11.py - TEPAT N11: menilai lembar jawaban review pakar buta dengan eval-kit
(hkeval.review.score_review).

Jalankan dari folder utama repo SETELAH lembar jawaban kembali dari pakar:
    python bukti/nilai_n11.py --isian ../review_n11/kembali/lembar_jawaban.xlsx --kit ../healthkathon-eval-kit

Membaca:
    --isian                       lembar_jawaban.xlsx yang sudah diisi pakar
    --key                         review_key.csv (bawaan ../review_n11/review_key.csv; DI LUAR repo publik)
    --sha                         bukti/hasil/n11_sha256.txt, sha256 key yang dicatat sebelum lembar dikirim
Menulis ke --out (bawaan bukti/hasil/):
    n11_review.csv                satu baris per kasus: review_id, kelompok, jawaban asli, nilai untuk kit
    ringkasan_n11.md              hasil score_review + penilai + catatan

Pemetaan ke nilai eval-kit (hkeval.review._normalise mengosongkan teks selain tiga nilai ini):
    Mencurigakan -> suspicious, Tidak mencurigakan -> not suspicious, Tidak bisa menilai -> cannot judge
    Kasus tidak realistis? Ya -> unrealistic_case = isi Catatan (atau "ya" bila Catatan kosong); Tidak/kosong -> ""

Skrip ini tidak membaca label maupun kit_truth: score_review hanya membandingkan porsi
"mencurigakan" di 20 kasus skor teratas (flagged) dan 20 kasus acak (random).
Butuh pandas, openpyxl, dan scipy (dipakai eval-kit).
"""
import argparse
import hashlib
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd

AKAR = Path(__file__).resolve().parent.parent
HASIL = AKAR / "bukti" / "hasil"
PETA_NILAI = {"mencurigakan": "suspicious", "tidak mencurigakan": "not suspicious",
              "tidak bisa menilai": "cannot judge"}
PETA_YT = {"ya": "Ya", "tidak": "Tidak"}
NAMA_KELOMPOK = {"flagged": "20 skor teratas (flagged)", "random": "20 acak dari sisanya (random)"}
BATAS_ID = ("Kecurigaan pakar bukan kecurangan yang terbukti. Pada kasus sintetis, ini menguji apakah pilihan "
            "sistem dan kasus buatan generator tampak masuk akal di mata pakar, bukan apakah sistem menemukan "
            "kecurangan nyata.")


def bersih(x):
    return re.sub(r"\s+", " ", str(x)).strip() if x is not None and str(x).strip() != "" else ""


def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def sha_tercatat(path, nama):
    for baris in Path(path).read_text(encoding="utf-8").splitlines():
        if baris and not baris.startswith("#"):
            h, f = baris.split(None, 1)
            if f.strip() == nama:
                return h
    raise SystemExit(f"{nama} tidak ada di {path}")


def baca_isian(path):
    """Kembalikan (DataFrame per kasus, dict info penilai) dari lembar_jawaban.xlsx."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    kepala = None
    for r in range(1, ws.max_row + 1):
        if bersih(ws.cell(r, 2).value) == "Kasus" and bersih(ws.cell(r, 3).value) == "Penilaian":
            kepala = r
            break
    if kepala is None:
        raise SystemExit("baris judul (Kasus, Penilaian) tidak ditemukan; apakah ini lembar_jawaban.xlsx N11?")
    kolom = [bersih(ws.cell(kepala, j).value) for j in range(1, 7)]
    assert kolom[:5] == ["No", "Kasus", "Penilaian", "Kasus tidak realistis?", "Catatan"], kolom
    info = {}
    for r in range(1, kepala):
        lab = bersih(ws.cell(r, 1).value)
        if lab and r > 2:
            info[lab] = bersih(ws.cell(r, 4).value)
    baris = []
    r = kepala + 1
    while bersih(ws.cell(r, 2).value).startswith("Kasus "):
        baris.append(dict(review_id=bersih(ws.cell(r, 2).value)[len("Kasus "):],
                          penilaian=bersih(ws.cell(r, 3).value),
                          tidak_realistis=bersih(ws.cell(r, 4).value),
                          catatan=bersih(ws.cell(r, 5).value)))
        r += 1
    return pd.DataFrame(baris), info


def petakan(isian):
    d = isian.copy()
    d["judgement"] = d["penilaian"].str.lower().map(PETA_NILAI).fillna("")
    d["tidak_realistis"] = d["tidak_realistis"].map(lambda x: PETA_YT.get(x.lower(), x) if x else "")
    d["unrealistic_case"] = [(c or "ya") if t == "Ya" else "" for t, c in zip(d["tidak_realistis"], d["catatan"])]
    aneh_nilai = d.loc[(d["penilaian"] != "") & (d["judgement"] == ""), "review_id"].tolist()
    aneh_yt = d.loc[~d["tidak_realistis"].isin(["", "Ya", "Tidak"]), "review_id"].tolist()
    return d, aneh_nilai, aneh_yt


def pct(x):
    return "-" if x != x else f"{x:.0%}"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--isian", required=True, help="lembar_jawaban.xlsx yang sudah diisi")
    ap.add_argument("--key", default=str(AKAR.parent / "review_n11" / "review_key.csv"))
    ap.add_argument("--sha", default=str(HASIL / "n11_sha256.txt"))
    ap.add_argument("--kit", default=str(AKAR.parent / "healthkathon-eval-kit"), help="folder eval-kit")
    ap.add_argument("--out", default=str(HASIL))
    a = ap.parse_args()
    sys.path.insert(0, str(Path(a.kit).resolve()))
    from hkeval.review import format_review, score_review

    # 1. key harus sama dengan yang dicatat sebelum lembar dikirim
    h_key, h_catat = sha256(a.key), sha_tercatat(a.sha, "review_key.csv")
    if h_key != h_catat:
        raise SystemExit(f"STOP: sha256 review_key.csv {h_key[:12]}... beda dengan yang dicatat {h_catat[:12]}... "
                         f"di {a.sha}. Key berubah setelah dicatat; jangan dinilai.")
    key = pd.read_csv(a.key)

    # 2. lembar jawaban
    isian, info = baca_isian(a.isian)
    if list(isian["review_id"]) != list(key["review_id"]):
        raise SystemExit("STOP: daftar/urutan kasus di lembar jawaban beda dengan review_key.csv")
    d, aneh_nilai, aneh_yt = petakan(isian)
    if aneh_nilai:
        print(f"PERINGATAN: Penilaian tidak dikenal (dihitung kosong): {', '.join(aneh_nilai)}")
    if aneh_yt:
        print(f"PERINGATAN: isian 'Kasus tidak realistis?' tidak dikenal: {', '.join(aneh_yt)}")

    # 3. eval-kit
    hasil = score_review(d[["review_id", "judgement", "unrealistic_case"]], key)
    teks_kit = format_review(hasil)

    # 4. tulis
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    rinci = d.merge(key[["review_id", "selection"]], on="review_id", how="left", validate="one_to_one")
    rinci[["review_id", "selection", "penilaian", "judgement", "tidak_realistis", "unrealistic_case",
           "catatan"]].to_csv(out / "n11_review.csv", index=False)

    try:
        kit_commit = subprocess.run(["git", "-C", a.kit, "rev-parse", "--short", "HEAD"], capture_output=True,
                                    text=True, check=True).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        kit_commit = "tidak diketahui"
    syarat = info.get("Saya bukan anggota tim, bukan juri atau panitia Healthkathon 2026", "")
    L = ["# Hasil review pakar buta (N11)", "",
         f"Lembar jawaban: `{Path(a.isian).name}` (sha256 {sha256(a.isian)[:16]}...). "
         f"Key cocok dengan sha256 yang dicatat sebelum lembar dikirim ({h_key[:16]}...). "
         f"Eval-kit commit `{kit_commit}`, `hkeval.review.score_review`.", "",
         "## Penilai", ""]
    L += [f"- {k}: {v or '(kosong)'}" for k, v in info.items()]
    if syarat != "Ya":
        L.append("- **PERHATIAN: syarat pakar (bukan tim/juri/panitia) tidak dikonfirmasi \"Ya\".**")
    L += ["", "## Porsi \"mencurigakan\" (suspicious / (suspicious + not suspicious))", "",
          "| Kelompok | Kasus | Dinilai | Mencurigakan | Porsi | IK95% Wilson | Tidak bisa menilai | Kosong |",
          "|---|---|---|---|---|---|---|---|"]
    for sel in ("flagged", "random"):
        g = hasil[sel]
        L.append(f"| {NAMA_KELOMPOK[sel]} | {g['cases']} | {g['judged']} | {g['suspicious']} | {pct(g['rate'])} | "
                 f"{pct(g['ci_low'])}-{pct(g['ci_high'])} | {g['cannot_judge']} | {g['blank']} |")
    dif = hasil["difference"]
    L += ["", f"Selisih (flagged dikurangi random): "
              f"{'-' if dif != dif else f'{dif * 100:+.0f} poin persen'}.",
          "Dengan 20 kasus per kelompok, selang Wilson lebar (sekitar plus-minus 20 poin); "
          "baca selisihnya bersama selangnya.", "",
          "## Kasus yang dinilai tidak realistis", ""]
    tr = rinci[rinci["unrealistic_case"] != ""]
    L += [f"- {r.review_id}: {r.unrealistic_case}" for r in tr.itertuples()] or ["- tidak ada"]
    lain = rinci[(rinci["catatan"] != "") & (rinci["unrealistic_case"] == "")]
    if len(lain):
        L += ["", "## Catatan lain dari pakar", ""] + [f"- {r.review_id}: {r.catatan}" for r in lain.itertuples()]
    if aneh_nilai or aneh_yt:
        L += ["", "## Isian tidak dikenal", "", f"- Penilaian: {', '.join(aneh_nilai) or '-'}",
              f"- Kasus tidak realistis?: {', '.join(aneh_yt) or '-'}"]
    L += ["", "## Batas", "", BATAS_ID, "", "Keluaran eval-kit apa adanya:", "", "```", teks_kit, "```", ""]
    (out / "ringkasan_n11.md").write_text("\n".join(L), encoding="utf-8")
    print(teks_kit)
    print(f"\nditulis: {out / 'n11_review.csv'}, {out / 'ringkasan_n11.md'}")


if __name__ == "__main__":
    main()
